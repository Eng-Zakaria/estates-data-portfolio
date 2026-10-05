"""Fetch DESNZ/DEFRA GHG conversion factors (flat workbook) and emit a dbt seed CSV.

Source: https://www.gov.uk/government/publications/greenhouse-gas-reporting-conversion-factors-2025
Workbook (flat format, for automatic processing):
https://assets.publishing.service.gov.uk/media/6846b6ea57f3515d9611f0dd/ghg-conversion-factors-2025-flat-format.xlsx
Licence: Open Government Licence v3.0.

Extracts Scope 1 natural-gas (kWh Net CV) and Scope 2 UK electricity (kWh)
kgCO2e factors by searching the 'Factors by Category' sheet — never hard-coded.
On failure writes seed_emission_factors_dev_placeholder.csv and exits non-zero.
"""

from __future__ import annotations

import csv
import sys
import urllib.request
from datetime import date
from pathlib import Path

URL = "https://assets.publishing.service.gov.uk/media/6846b6ea57f3515d9611f0dd/ghg-conversion-factors-2025-flat-format.xlsx"
YEAR = 2025
SEED = Path(__file__).resolve().parents[1] / "seeds" / "seed_emission_factors.csv"
PLACEHOLDER = Path(__file__).resolve().parents[1] / "seeds" / "seed_emission_factors_dev_placeholder.csv"

WANTED = [
    # (fuel, scope, level1, level2, level3, uom)
    ("natural_gas", "Scope 1", "Fuels", "Gaseous fuels", "Natural gas", "kWh (Net CV)"),
    ("electricity", "Scope 2", "UK electricity", None, None, "kWh"),
]


def fetch(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    urllib.request.urlretrieve(URL, path)
    print(f"downloaded {URL} -> {path}")


def extract(workbook: Path) -> list[dict[str, object]]:
    import openpyxl

    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    ws = wb["Factors by Category"]
    header_idx = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if row[0] == "ID":
            header_idx = i
            break
    if header_idx is None:
        raise ValueError("header row (ID, Scope, ...) not found")
    rows = list(ws.iter_rows(min_row=header_idx + 2, values_only=True))
    found: list[dict[str, object]] = []
    for fuel, scope, l1, l2, l3, uom in WANTED:
        match = None
        for r in rows:
            _id, sc, a, b, c, _d, _ct, u, ghg, factor = r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9]
            if sc != scope or u != uom or ghg != "kg CO2e":
                continue
            if a != l1:
                continue
            if l2 is not None and b != l2:
                continue
            if l3 is not None and c != l3:
                continue
            match = (r[0], float(factor))
            break
        if match is None:
            raise ValueError(f"no factor row for {fuel} {scope} {uom}")
        found.append(
            {
                "fuel": fuel,
                "scope": scope,
                "unit": "kWh",
                "factor_kgco2e_per_kwh": match[1],
                "year": YEAR,
                "source_url": URL,
            }
        )
    return found


def main() -> int:
    cache = Path("data/reference/desnz-flat-2025.xlsx")
    try:
        fetch(cache)
        rows = extract(cache)
    except Exception as exc:  # noqa: BLE001 — loud fallback per hard rules
        print(f"DESNZ fetch/extract failed: {exc}", file=sys.stderr)
        PLACEHOLDER.parent.mkdir(parents=True, exist_ok=True)
        with open(PLACEHOLDER, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=["fuel", "scope", "unit", "factor_kgco2e_per_kwh", "year", "source_url"])
            w.writeheader()
            w.writerow(
                {
                    "fuel": "DEV_PLACEHOLDER",
                    "scope": "Scope 1",
                    "unit": "kWh",
                    "factor_kgco2e_per_kwh": 0.0,
                    "year": YEAR,
                    "source_url": URL,
                }
            )
        print(f"wrote NON-AUTHORITATIVE placeholder {PLACEHOLDER}; see docs/OPEN_QUESTIONS.md", file=sys.stderr)
        return 1
    SEED.parent.mkdir(parents=True, exist_ok=True)
    with open(SEED, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["fuel", "scope", "unit", "factor_kgco2e_per_kwh", "year", "source_url"])
        w.writeheader()
        w.writerows(rows)
    for r in rows:
        print(f"{r['fuel']}: {r['factor_kgco2e_per_kwh']} kgCO2e/kWh ({r['scope']}, {YEAR})")
    print(f"retrieved {date.today().isoformat()}; wrote {SEED}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
